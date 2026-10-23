"""Imported local reference assets, independent of the engineer's open workbook."""
import json
from pathlib import Path
import tempfile
import uuid

from .templates import sha256


def reference_file(registry_path, role='aa', source=None):
    """Import once; only an explicit source argument refreshes the local asset."""
    from .connections import copy_source
    registry_path=Path(registry_path)
    config=json.loads(registry_path.read_text(encoding='utf8'))
    key='runtime_file' if role=='aa' else 'wall_runtime_file'
    existing=config.get(key)
    if source is None and existing and Path(existing).is_file():
        return Path(existing)
    selected=Path(source or (config.get('connection_file',config.get('blinding_file')) if role=='aa' else config['file']))
    folder=registry_path.parent/'runtime';folder.mkdir(exist_ok=True)
    target=folder/(('aa-' if role=='aa' else 'wall-')+uuid.uuid4().hex[:12]+selected.suffix.lower())
    try:
        copy_source(selected,target)
        if role=='aa':
            from openpyxl import load_workbook
            book=load_workbook(target,read_only=True,data_only=False,keep_links=False)
            try:
                required=('Connection check','Bracket','Blinding','CS1 ','W1','D-Strut&Wailing Summary','Section Table UC')
                missing=[name for name in required if name not in book.sheetnames]
                if missing:raise ValueError('参考文件缺少工作表：'+', '.join(missing))
            finally:book.close()
        else:
            import xlrd
            book=xlrd.open_workbook(target)
            if 'FSPiv Summary' not in book.sheet_names():raise ValueError('参考文件缺少钢板桩验算页')
        config[key]=str(target.resolve())
        config['reference_source' if role=='aa' else 'wall_reference_source']=str(selected.resolve())
        config['reference_revision' if role=='aa' else 'wall_reference_revision']=sha256(target)
        staged=registry_path.with_suffix('.pending.json')
        staged.write_text(json.dumps(config,ensure_ascii=False,indent=2),encoding='utf8');staged.replace(registry_path)
        return target
    except Exception:
        target.unlink(missing_ok=True)
        raise


def detach_external_links(wb, printed_sheets):
    """Keep generated formula dependencies local and discard historical Excel links."""
    before={name:wb.Worksheets(name).Range(wb.Worksheets(name).PageSetup.PrintArea).Value2 for name in printed_sheets}
    links=wb.LinkSources(1)
    if isinstance(links,str):links=(links,)
    for link in links or ():
        wb.BreakLink(link,1)
    for index in range(wb.Names.Count,0,-1):
        name=wb.Names.Item(index)
        if '[' in str(name.RefersTo):
            name.Delete()
    if wb.LinkSources(1):raise ValueError('生成副本仍包含外部工作簿链接')
    for name,values in before.items():
        ws=wb.Worksheets(name)
        if ws.Range(ws.PageSetup.PrintArea).Value2!=values:
            raise ValueError(f'{name}包含尚未本地化的外部计算依赖')
