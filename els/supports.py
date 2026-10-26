"""Template-driven steel members, direct fillet welds and blinding concrete."""
import math
import re
from pathlib import Path
import json
import copy

from .core import finite
from .templates import sha256
from .excel_ops import batched_print, generation_run, progress, finish_headers, commit_print

ROOT = Path(__file__).resolve().parent.parent
BLINDING_FILE = ROOT / '新计算书/Structural capacity AA section.xlsx'
MEMBER_FIELDS = ('member_id', 'case', 'kind', 'section', 'grade', 'anchor_kn_per_m',
                 'load_width_m', 'load_span_m', 'angle_deg', 'lx_m', 'ly_m', 'kx', 'ky', 'klt')
WELD_FIELDS = ('member_id', 'case', 'kind', 'leg_mm', 'strength_mpa', 'effective_length_mm', 'distribution_length_m')


def migrate_package(package):
    """Keep legacy design actions while converting the displayed force to kN."""
    result=copy.deepcopy(package or {})
    if result.get('schema_version')==2:
        return result
    members=result.get('members',[])
    struts=[m for m in members if m.get('kind')=='Strut']
    for m in struts:
        m['anchor_kn']=abs(m['anchor_kn_per_m'])*m['load_width_m'] if finite(m.get('anchor_kn_per_m')) and finite(m.get('load_width_m')) else None
        m['spacing_m']=m.get('load_width_m')
        m['strut_length_m']=m.get('load_span_m')
        m['lx_is_default']=False
        m['lx_equals_ly']=False
    for m in members:
        if m.get('kind')!='Waler':continue
        matches=[s for s in struts if s.get('case')==m.get('case') and
                 s.get('spacing_m')==m.get('load_width_m') and
                 s.get('anchor_kn_per_m')==m.get('anchor_kn_per_m')]
        if len(matches)==1:
            m['source_strut_id']=matches[0]['member_id'];m['source_strut_case']=matches[0]['case']
        else:
            m['source_strut_id']='';m['source_strut_case']=''
        m['lx_equals_ly']=False
    result['schema_version']=2
    return result


def resolve_package(package):
    """Normal force kN -> one strut axial force; linked waler q = force/spacing."""
    result=copy.deepcopy(package)
    if result.get('schema_version')!=2:return result
    linked={(m.get('member_id'),m.get('case')):m for m in result.get('members',[]) if m.get('kind')=='Strut'}
    for m in result.get('members',[]):
        if m.get('kind')=='Strut':
            m['load_width_m']=m.get('spacing_m');m['load_span_m']=m.get('strut_length_m')
            m['anchor_kn_per_m']=abs(m['anchor_kn'])/m['spacing_m'] if finite(m.get('anchor_kn')) and finite(m.get('spacing_m')) and m['spacing_m']>0 else None
            if m.get('lx_is_default'):m['lx_m']=m.get('strut_length_m')
        else:
            source=linked.get((m.get('source_strut_id'),m.get('source_strut_case')))
            spacing=source.get('spacing_m') if source else None
            force=source.get('anchor_kn') if source else None
            m['load_width_m']=spacing;m['angle_deg']=0.0
            m['anchor_kn_per_m']=abs(force)/spacing if finite(force) and finite(spacing) and spacing>0 else None
            if m.get('lx_is_default'):m['lx_m']=m.get('load_span_m')
        if m.get('lx_equals_ly'):m['ly_m']=m.get('lx_m')
        if m.get('use_default_k'):
            m.update(kx=1.0,ky=1.0,klt=1.0)
    for bracket in result.get('brackets',[]):
        bracket['eccentricity_m']=bracket['length_m']/2 if finite(bracket.get('length_m')) else None
    from .critical_connections import expand_groups
    expand_groups(result)
    return result


def section_library(aa=False):
    if aa:
        from tempfile import TemporaryDirectory
        from openpyxl import load_workbook
        from .connections import copy_source
        from .reference_store import reference_file
        imported=reference_file(ROOT/'templates/supports-structure.local.json')
        with TemporaryDirectory(prefix='els-aa-sections-') as folder:
            snapshot=Path(folder)/'aa.xlsx';copy_source(imported,snapshot)
            book=load_workbook(snapshot,read_only=True,data_only=True,keep_links=False)
            try:
                return [r[0] for r in book['Section Table UC'].iter_rows(min_row=5,max_row=128,max_col=1,values_only=True) if isinstance(r[0],str) and 'x' in r[0]]
            finally:book.close()
    import xlrd
    config = json.loads((ROOT / 'templates/sheetpile-structure.local.json').read_text(encoding='utf-8'))
    b = xlrd.open_workbook(config['file'])
    s = b.sheet_by_name('Section Table UC')
    return [s.cell_value(row, 0) for row in range(4, min(128, s.nrows))
            if s.cell_type(row, 0) == 1 and 'x' in s.cell_value(row, 0)]


def validate_package(package, library=None):
    if not isinstance(package, dict):
        return ['构件输入须为对象']
    package=resolve_package(package)
    errors, seen = [], set()
    from .critical_connections import validate_groups
    errors.extend(validate_groups(package))
    members = package.get('members', [])
    if not isinstance(members, list) or (not members and not package.get('include_blinding')):
        return ['至少录入一个支撑或围檩工况']
    for i, m in enumerate(members, 1):
        if not isinstance(m, dict):
            errors.append(f'构件{i}格式无效')
            continue
        key = (m.get('member_id'), m.get('case'))
        if any(not isinstance(x, str) or not re.fullmatch(r'[A-Za-z0-9_-]{1,12}', x) for x in key):
            errors.append(f'构件{i}编号及工况须为1～12位英文字母／数字／短横线')
        elif key in seen:
            errors.append(f'构件{i}编号与工况重复')
        else:
            seen.add(key)
        if m.get('kind') not in ('Strut', 'Waler'):
            errors.append(f'构件{i}类型无效')
        if not m.get('section') or (library is not None and m['section'] not in library):
            errors.append(f'构件{i}截面不在模板库')
        if m.get('grade') not in ('S235', 'S275', 'S355', 'S460'):
            errors.append(f'构件{i}钢材等级无效')
        if package.get('schema_version')==2 and m.get('kind')=='Strut' and not finite(m.get('anchor_kn')):
            errors.append(f'构件{i}墙法线力须为有效kN数值')
        if package.get('schema_version')==2 and m.get('kind')=='Waler' and not finite(m.get('load_width_m')):
            errors.append(f'围檩{i}请关联已录入的支撑及工况；间距由支撑自动提供')
        for field in MEMBER_FIELDS[5:]:
            v = m.get(field)
            if not finite(v) or (field not in ('anchor_kn_per_m', 'angle_deg') and v <= 0):
                label = {'load_width_m': '支撑间距 m', 'load_span_m': 'Strut length m' if m.get('kind') == 'Strut' else 'Span m'}.get(field, field)
                errors.append(f'构件{i} {label}须填写有效数值')
        if finite(m.get('angle_deg')) and not 0 <= m['angle_deg'] < 90:
            errors.append(f'构件{i}角度须在0～90度之间，不含90度')
    welds = package.get('welds', [])
    if not isinstance(welds, list) or (not welds and package.get('schema_version')!=2):
        errors.append('至少录入一项模板直接受力焊缝连接')
        welds = []
    linked = {(m.get('member_id'), m.get('case')): m for m in members if isinstance(m, dict)}
    weld_seen = set()
    for i, w in enumerate(welds, 1):
        if not isinstance(w, dict):
            errors.append(f'焊缝{i}格式无效')
            continue
        key = (w.get('member_id'), w.get('case'))
        m = linked.get(key)
        kind = w.get('kind')
        if kind=='Strut-Waler' and not w.get('group_id'):
            from .critical_connections import associated_walers
            if not associated_walers(package,w):errors.append(f'焊缝{i}缺少对应围檩，需关联围檩以读取设计剪力')
        if not m or kind not in ('Strut-Waler', 'Waler-Sheetpile') or m.get('kind') != ('Strut' if kind == 'Strut-Waler' else 'Waler'):
            errors.append(f'焊缝{i}构件关联或连接类型错误')
        if key + (kind,) in weld_seen:
            errors.append(f'焊缝{i}连接重复')
        weld_seen.add(key + (kind,))
        for field in ('leg_mm', 'strength_mpa') + (('effective_length_mm', 'distribution_length_m') if kind == 'Waler-Sheetpile' else ()):
            if not finite(w.get(field)) or w[field] <= 0:
                errors.append(f'焊缝{i} {field}须填写正数')
    for key in (('length_m', 'width_m', 'blinding_thickness_mm') if package.get('include_blinding',True) else ()):
        if not finite(package.get(key)) or package[key] <= 0:
            errors.append(f'{key}须填写正数；垫层有效长度取基坑短边')
    from .connections import AA_CONNECTIONS, validate_brackets
    if package.get('connection_template') not in (None, AA_CONNECTIONS):
        errors.append('不支持的连接模板版本')
    if package.get('brackets') and package.get('connection_template') != AA_CONNECTIONS:
        errors.append('托架须使用AA连接模板')
    errors.extend(validate_brackets(package, library))
    if not isinstance(package.get('source'), str) or not package['source'].strip():
        errors.append('请填写anchor force分析来源及版本')
    return errors


def member_actions(m, area_cm2):
    w = abs(m['anchor_kn_per_m'])
    q = 1.4 * 77 * area_cm2 / 10000
    if m['kind'] == 'Strut':
        normal_force=abs(m['anchor_kn']) if 'anchor_kn' in m else w*m['load_width_m']
        return {'n': 1.4*normal_force/math.cos(math.radians(m['angle_deg'])),
                'v': q*m['load_span_m']/2, 'mx': q*m['load_span_m']**2/8, 'my': 0.0}
    return {'n': 1.4*w*m['load_width_m'], 'v': 1.4*w*m['load_span_m']/2,
            'mx': 1.4*w*m['load_span_m']**2/8, 'my': q*m['ly_m']**2/8}


def blinding_reference(force_kn_per_m, length_m, width_m, thickness_mm):
    le = min(length_m, width_m)*1000
    h = thickness_mm
    ex = max(h/20, 20)
    ea = le**2/(2500*h)
    capacities = [0.3*(h-2*ex)*20, 0.3*(h-2*(ex+ea))*20]
    demand = 1.4*abs(force_kn_per_m)
    return {'effective_length_mm': le, 'thickness_mm': h, 'fcu_mpa': 20,
            'demand_kn_per_m': demand, 'eccentricity_mm': ex, 'additional_eccentricity_mm': ea,
            'capacities_kn_per_m': capacities, 'slenderness': le/h,
            'status': 'PASS' if le/h <= 30 and min(capacities) >= demand and min(capacities)>0 else 'FAIL'}


def numeric(ws, address):
    value = ws.Range(address).Value2
    if not finite(value) or str(ws.Range(address).Text) in ('#N/A', '#REF!', '#VALUE!', '#DIV/0!', '#NUM!', '#NAME?'):
        raise ValueError(f'{ws.Name}!{address}公式未返回有效数值：{ws.Range(address).Text}，公式 {ws.Range(address).Formula}')
    return float(value)


@batched_print
def page_style(ws, project, subject, last_row, app):
    ws.Visible = -1
    ws.Range('D3').Value2 = "'" + project
    ws.Range('D4').Value2 = "'" + subject
    ws.Range('Q4').NumberFormat = '[$-409]dd-mmm-yy'
    ws.Range('Q4').ShrinkToFit = True
    ws.Range('Q4').Font.Size = 8
    for address in ('R3', 'V3', 'AB3', 'AG3', 'V4', 'AB4'):
        ws.Range(address).MergeArea.ClearContents()
    for index in range(ws.Shapes.Count, 0, -1):
        shape = ws.Shapes.Item(index)
        if shape.Type in (11, 13) and (shape.TopLeftCell.Row <= 2 or shape.TopLeftCell.Row > last_row):
            shape.Delete()
    header = ws.Range('A2:AH2')
    header.UnMerge()
    header.Merge()
    header.Value2 = 'STRUCTURAL CALCULATIONS'
    header.Font.Name = 'Arial'
    header.Font.Bold = True
    header.HorizontalAlignment = -4108
    from .printing import setup_print
    setup_print(ws,last_row,app)
    ws.PageSetup.CenterHeader = 'REVIEW DRAFT'
    ws.PageSetup.CenterFooter = ''


def report_row(ws, row, values, ranges, header=False):
    numeric_ranges=[]
    for value, (left, right) in zip(values, ranges):
        cells=ws.Range(f'{left}{row}:{right}{row}')
        cells.Merge()
        cells.Value2=value
        if finite(value):
            numeric_ranges.append(f'{left}{row}:{right}{row}')
    area=ws.Range(','.join(f'{left}{row}:{right}{row}' for left,right in ranges))
    area.WrapText=True
    font=area.Font;font.Size=9;font.Bold=header
    area.Borders.LineStyle=1
    area.VerticalAlignment=-4108
    if numeric_ranges:ws.Range(','.join(numeric_ranges)).NumberFormat='0.000'
    ws.Rows(f'{row}:{row}').RowHeight=32


def populate(wb, package, actions, project, subject):
    progress('members','生成支撑、围檩及构件验算…')
    package=resolve_package(package)
    from .connections import AA_CONNECTIONS
    errors = validate_package(package, section_library(aa=package.get('connection_template')==AA_CONNECTIONS))
    if errors:
        raise ValueError('\n'.join(errors))
    app = wb.Application
    include_blinding=package.get('include_blinding',True)
    if include_blinding:
        registry = json.loads((ROOT/'templates/supports-structure.local.json').read_text(encoding='utf-8'))
        from .reference_store import reference_file
        blinding_file=reference_file(ROOT/'templates/supports-structure.local.json')
    output, results, member_sheets = [], [], {}
    def copy_sheet(name):
        source = wb.Worksheets(name)
        visibility = source.Visible
        source.Visible = -1
        try:
            source.Copy(wb.Worksheets('FSPiv Summary'))
            if app.ActiveSheet.Parent.FullName != wb.FullName:
                raise ValueError('Excel模板页未复制到目标工作簿')
            return app.ActiveSheet
        finally:
            source.Visible = visibility
    summary = copy_sheet('Strut&Wailing Summary')
    summary.Name = 'Strut & Wailing Summary'
    summary.Range('A7:AH100').UnMerge()
    summary.Range('A7:AH100').Clear()
    inputs=copy_sheet('Strut&Wailing Summary')
    inputs.Name='Member Data'
    inputs.Range('A7:AH100').UnMerge()
    inputs.Range('A7:AH100').Clear()
    input_spans=[('A','B'),('C','D'),('E','F'),('G','J'),('K','L'),('M','O'),('P','R'),('S','U'),('V','W'),('X','Y'),('Z','AA'),('AB','AC'),('AD','AE'),('AF','AH')]
    report_row(inputs,8,('Member','Case','Type','Section','Grade','Normal kN / waler kN/m','Spacing m','Length / Span m','Angle deg','Lx m','Ly m','Kx','Ky','KLT'),input_spans,True)
    for index, m in enumerate(package['members'], 9):
        values=[m[k] for k in MEMBER_FIELDS]
        if m['kind']=='Strut' and 'anchor_kn' in m:values[5]=m['anchor_kn']
        report_row(inputs,index,values,input_spans)
        source = 'CCS1 ' if m['kind'] == 'Strut' else 'CW1'
        ws = copy_sheet(source)
        ws.Name = f"{'S' if m['kind']=='Strut' else 'W'}-{m['member_id']}-{m['case']}"
        # Resolve table lookups inside this workbook rather than copied external names.
        ws.Cells.Replace(What=',Table,',Replacement=",'Section Table UC'!$A$5:$X$128,",
                         LookAt=2,SearchOrder=1,MatchCase=False,SearchFormat=False,ReplaceFormat=False)
        ref=lambda col: f"'Member Data'!{col}{index}"
        ws.Range('P11').Formula = '='+ref('G')
        ws.Range('Y11').Formula = '='+ref('K')
        ws.Range('F8').Value2 = m['member_id']+' / '+m['case']
        for cell, col in [('J16','X'), ('J18','Z'), ('G17','AB'), ('G19','AD'), ('G20','AF')]:
            ws.Range(cell).Formula = '='+ref(col)
        ws.Calculate()
        loads = member_actions(m, numeric(ws,'R15'))
        ws.Range('J11').Formula=f'=1.4*ABS({ref("M")})' + ('' if m['kind']=='Strut' and 'anchor_kn' in m else f'*{ref("P")}') + (f'/COS(RADIANS({ref("V")}))' if m['kind']=='Strut' else '')
        action_formula=ws.Range('J11').Formula[1:]
        required=','.join(ref(c) for c in ('M','P','S','V','X','Z','AB','AD','AF'))
        positive=','.join(ref(c) for c in ('P','S','X','Z','AB','AD','AF'))
        ws.Range('J11').Formula=f'=IF(AND(COUNT({required})=9,MIN({positive})>0,{ref("V")}>=0,{ref("V")}<90),{action_formula},"INPUTS REQUIRED")'
        if m['kind']=='Strut':
            ws.Range('J12').Formula=f'=1.4*77*R15/10000*{ref("S")}/2'
            ws.Range('J13').Formula=f'=J12*{ref("S")}/4'
            ws.Range('J14').Value2=0
        else:
            ws.Range('J12').Formula=f'=1.4*ABS({ref("M")})*{ref("S")}/2'
            ws.Range('J13').Formula=f'=1.4*ABS({ref("M")})*{ref("S")}^2/8'
            ws.Range('J14').Formula='=1.4*77*R15/10000*J18^2/8'
        ws.Calculate()
        for cell,key in [('J11','n'),('J12','v'),('J13','mx'),('J14','my')]:
            if not math.isclose(numeric(ws,cell),loads[key],rel_tol=1e-9,abs_tol=1e-9):
                raise ValueError(f'{ws.Name}!{cell}内力公式与独立算式不一致')
        # Preserve the engineer-corrected template references; reject stale copies.
        for cell in ('L46', 'L48', 'L50'):
            if not math.isclose(numeric(ws, cell), loads['n'], rel_tol=1e-9, abs_tol=1e-9):
                raise ValueError(f'{ws.Name}!{cell}屈曲轴力未引用本构件设计轴力')
        checks = {cell: numeric(ws, cell) for cell in ('U39','U46','U48','U50')}
        checks['shear'] = loads['v']/numeric(ws,'E29')
        checks['compression'] = loads['n']/min(numeric(ws,'U35'),numeric(ws,'U36'))
        result = {'member': m, 'design_actions': loads, 'checks': checks,
                  'status': 'PASS' if max(checks.values())<=1 else 'FAIL',
                  'sheet': ws.Name, 'yield_strength_mpa': numeric(ws,'AD11')}
        results.append(result)
        member_sheets[(m['member_id'],m['case'])] = (ws, result)
        page_style(ws,project,subject,51,app)
        output.append(ws.Name)
    from .support_layout import build_summary, build_welds
    progress('member_summary','生成构件总体验算表…')
    summary_last=build_summary(wb,summary,package,member_sheets)
    page_style(summary,project,subject,summary_last,app)
    inputs.Visible=2
    if package['members']:output.insert(0,summary.Name)
    else:summary.Visible=2
    weld_results=[]
    bracket_results=[]
    connection_hash=None
    if package.get('welds') or package.get('brackets'):
        progress('connections','生成控制连接及托架验算…')
        weld_sheet=copy_sheet('Connection check')
        weld_sheet.Name='Connection Checks'
        from .connections import AA_CONNECTIONS, build_aa_connections
        if package.get('connection_template') == AA_CONNECTIONS:
            weld_results,bracket_results,weld_last,connection_hash=build_aa_connections(wb,weld_sheet,package,member_sheets,numeric)
        else:
            weld_results,weld_last=build_welds(wb,weld_sheet,package,member_sheets,numeric)
        page_style(weld_sheet,project,subject,weld_last,app)
        output.append(weld_sheet.Name)
        for bracket in bracket_results:
            detail=wb.Worksheets(bracket['calculation_sheet'])
            detail.Visible=-1
            page_style(detail,project,subject,51,app)
            output.append(detail.Name)
    expected=None
    source_hash=None
    if include_blinding:
        progress('blinding','生成 Blinding concrete 验算…')
        # Copy only the engineer-selected Blinding worksheet from the separate source.
        from .connections import controlled_source
        with controlled_source(app,blinding_file,sha256(blinding_file)) as (other,source_hash):
            # XLS has 65,536 rows: copy the controlled region of the XLSX source.
            blind=copy_sheet('Blinding')
            blind.Range('A1:AH100').UnMerge()
            blind.Range('A1:AH100').Clear()
            source_blind=other.Worksheets('Blinding')
            for address,expected_formula in registry['expected_formulas'].items():
                if source_blind.Range(address).Formula != expected_formula:
                    raise ValueError(f'Blinding!{address}公式与登记版本不同')
            source_blind.Range('A1:AH27').Copy(blind.Range('A1'))
            for col in range(1,35):
                blind.Columns(col).ColumnWidth=source_blind.Columns(col).ColumnWidth
            for row in range(1,28):
                blind.Rows(row).RowHeight=source_blind.Rows(row).RowHeight
        blind.Name='Blinding Concrete'
        blind.Range('D8').Formula="='FSPiv Summary'!F50"
        blind.Range('H9').Value2=1.4
        blind.Range('D10').Value2=20
        blind.Range('D11').Value2=package['blinding_thickness_mm']
        blind.Range('D12').Value2=min(package['length_m'],package['width_m'])*1000
        blind.Range('F13').Formula='=IF(D13<=G13,"<=",">")'
        blind.Range('H13').Formula='=IF(AND(COUNT(D11,D12)=2,MIN(D11,D12)>0),IF(D13<=G13,"PASS","FAIL"),"INPUTS REQUIRED")'
        for row in (22,26):
            blind.Range(f'F{row}').Formula=f'=IF(C{row}>=G{row},">=","<")'
            blind.Range(f'I{row}').Formula=f'=IF(AND(C{row}>0,C{row}>=G{row}),"PASS","FAIL")'
            blind.Range(f'E{row}').Value2='kN/m'
        blind.Range('B7').Value2='Assumption: blinding axial force equals the maximum unfactored sheetpile shear; effective length equals the pit short side.'
        blind.Range('C24').Value2='0.3[h - 2(ex,2 + ea)]fcu'
        page_style(blind,project,subject,27,app)
        blind.Calculate()
        expected=blinding_reference(actions['v'],package['length_m'],package['width_m'],package['blinding_thickness_mm'])
        for cell,value in [('N9',expected['demand_kn_per_m']),('C22',expected['capacities_kn_per_m'][0]),('C26',expected['capacities_kn_per_m'][1])]:
            if not math.isclose(numeric(blind,cell),value,rel_tol=1e-9):
                raise ValueError('Blinding公式与独立算式不一致')
        output.append(blind.Name)
    basis = ('Unfactored normal strut force kN; axial force = 1.4F/cos(angle); linked waler q = F/spacing; '
             if package.get('schema_version') == 2 else
             'Unfactored anchor force kN/m; strut axial force = 1.4q*load width/cos(angle); ')
    result={'members':results,'welds':weld_results,'brackets':bracket_results,'connection_source_sha256':connection_hash,'blinding':expected,'blinding_source_sha256':source_hash,
            'assumptions':basis+'Template selfweight and concentric fillet welds.'}
    failures=[m['sheet'] for m in results if m['status']=='FAIL']+[w['input']['member_id']+' weld' for w in weld_results if w['status']=='FAIL']
    failures.extend(b['input']['bracket_id']+' bracket' for b in bracket_results if b['status']=='FAIL')
    if expected and expected['status']=='FAIL':
        failures.append('Blinding Concrete')
    result['failed_checks']=failures
    result['status']='NOT PASS' if failures else 'PASS'
    return output,result

def generate_supports(package, mapping_path, output_root, project_name, subject_name, steel_request=None, progress=None):
    with generation_run(progress) as run:
        return run.finish(_generate_supports(package,mapping_path,output_root,project_name,subject_name,steel_request))


def _generate_supports(package, mapping_path, output_root, project_name, subject_name, steel_request=None):
    """Generate member/connection checks without requiring a sheetpile design."""
    import hashlib
    import shutil
    import uuid
    from datetime import datetime
    from .sheetpile import excel_session, check_template, envelope
    from .connections import AA_CONNECTIONS
    errors=validate_package(package,section_library(aa=package.get('connection_template')==AA_CONNECTIONS))
    if errors:raise ValueError('\n'.join(errors))
    if not project_name.strip() or not subject_name.strip() or re.search(r'[\u3400-\u9fff]',project_name+subject_name):raise ValueError('请填写英文Project Name及Subject Name')
    actions={'v':0.0}
    if package.get('include_blinding',True):
        if not steel_request:raise ValueError('生成Blinding需先填写钢板桩剪力，或取消Blinding勾选')
        actions=envelope(steel_request)
    from .connections import AA_CONNECTIONS,copy_source,prepare_aa_workbook
    aa=package.get('connection_template')==AA_CONNECTIONS
    if aa:
        from .reference_store import reference_file
        source=reference_file(ROOT/'templates/supports-structure.local.json')
    else:
        mapping=json.loads(Path(mapping_path).read_text(encoding='utf8'))
        source=Path(mapping['file']);expected_hash=mapping['sha256']
    folder=Path(output_root)/(datetime.now().strftime('%Y%m%d-%H%M%S')+'-'+uuid.uuid4().hex[:6]);folder.mkdir(parents=True,exist_ok=False)
    target=folder/('Waling_Strut_Check_Review.xlsx' if aa else 'Waling_Strut_Check_Review.xls');pdf=folder/'Waling_Strut_Check_Review.pdf'
    try:
        if aa:copy_source(source,target)
        else:shutil.copy2(source,target)
        before=sha256(target)
        with excel_session(target,defer_print=True) as (app,wb):
            progress('prepare_template','准备计算书模板及截面库…')
            if aa:prepare_aa_workbook(wb)
            else:check_template(wb)
            main=wb.Worksheets('FSPiv Summary');main.Visible=-1
            for ws in wb.Worksheets:
                if ws.Name!=main.Name:ws.Visible=2
            main.Range('F50').Value2=actions['v']
            names,checks=populate(wb,package,actions,project_name,subject_name)
            main.Visible=2
            for name in reversed(names):
                if wb.Worksheets(1).Name!=name:wb.Worksheets(name).Move(wb.Worksheets(1))
            progress('finish_layout','整理页眉及打印设置…')
            finish_headers(wb,target.name,checks.get('status')=='NOT PASS')
            commit_print(app)
            progress('verify_formulas','重算并检查输出公式…')
            app.Calculation=-4105;app.CalculateFullRebuild()
            for name in names:
                ws=wb.Worksheets(name)
                try:bad=ws.Range(ws.PageSetup.PrintArea).SpecialCells(-4123,16)
                except Exception:bad=None
                if bad is not None:raise ValueError(f'{name}打印区域存在公式错误：{bad.Address}')
            from .reference_store import detach_external_links
            progress('verify_links','检查并清除外部链接，核对计算结果…')
            detach_external_links(wb,names)
            from .printing import format_check_results
            for name in names:format_check_results(wb.Worksheets(name))
            from .sheetpile import save_and_export
            save_and_export(wb,pdf.resolve())
        result={'module':'support_structure','request':package,'project_name':project_name,'subject_name':subject_name,
                'steel_request':steel_request if package.get('include_blinding',True) else None,
                'template_sha256':before,'output_sheets':names,'support_package_results':checks,
                'review_status':'ENGINEER_REVIEW_REQUIRED','workbook':str(target.resolve()),'pdf':str(pdf.resolve()),
                'input_sha256':hashlib.sha256(json.dumps({'package':package,'project_name':project_name,'subject_name':subject_name,'blinding_shear':actions['v']},sort_keys=True,ensure_ascii=False,allow_nan=False).encode('utf8')).hexdigest()}
        return result
    except Exception as exc:
        for file in (target,pdf):file.unlink(missing_ok=True)
        (folder/'failure.json').write_text(json.dumps({'error':str(exc),'request':package},ensure_ascii=False,indent=2),encoding='utf8');raise
