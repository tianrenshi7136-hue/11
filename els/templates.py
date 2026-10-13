"""Read-only template inspection and audited manifest checks."""
import hashlib
import json
import re
import zipfile
from pathlib import Path


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def inspect_excel(path):
    from openpyxl import load_workbook
    path = Path(path)
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("当前检查器支持xlsx/xlsm；xls需单独接入Excel检查")
    with zipfile.ZipFile(path) as archive:
        parts = archive.namelist()
        external_links = [p for p in parts if p.startswith("xl/externalLinks/") and p.endswith(".xml")]
        macros = "xl/vbaProject.bin" in parts
    book = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    try:
        sheets = []
        for sheet in book:
            formulas = [{"cell": c.coordinate, "formula": str(c.value)} for row in sheet for c in row if c.data_type == "f"]
            sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column, "state": sheet.sheet_state, "formula_count": len(formulas), "formula_sample": formulas[:30], "print_area": str(sheet.print_area or ""), "merged_cells": [str(r) for r in sheet.merged_cells.ranges]})
        return {"path": str(path.resolve()), "sha256": sha256(path), "macros": macros, "external_links": external_links, "defined_names": {k: str(v.attr_text) for k, v in book.defined_names.items()}, "sheets": sheets, "status": "仅检查结构，未计算、未审核"}
    finally:
        book.close()


def check_manifest(path):
    path = Path(path).resolve()
    manifest = json.loads(path.read_text(encoding="utf-8-sig"))
    problems = []
    if manifest.get("schema_version") != 1:
        problems.append("模板映射版本不支持")
    if manifest.get("approved") is not True or not manifest.get("reviewer") or not manifest.get("rules_reference"):
        problems.append("模板计算方法尚未完成工程审核")
    for kind, suffixes in (("excel", (".xlsx", ".xlsm")), ("cad", (".dwg",))):
        spec = manifest.get(kind, {})
        file = (path.parent / spec.get("file", "")).resolve()
        if not file.is_file() or file.suffix.lower() not in suffixes:
            problems.append(f"{kind}模板文件不存在或格式不支持")
        elif sha256(file) != spec.get("sha256"):
            problems.append(f"{kind}模板指纹不符，需重新审核")
    mappings = manifest.get("excel", {}).get("inputs", [])
    if not mappings:
        problems.append("尚未建立Excel输入映射")
    targets = set()
    for mapping in mappings:
        target = (mapping.get("sheet"), mapping.get("cell"))
        if not target[0] or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", target[1] or "") or not mapping.get("source"):
            problems.append("Excel映射缺少工作表、单元格或数据来源")
        if target in targets:
            problems.append("Excel输入映射包含重复单元格")
        targets.add(target)
    if not manifest.get("excel", {}).get("checks"):
        problems.append("尚未定义Excel验算结果及通过判定")
    if not manifest.get("cad", {}).get("objects"):
        problems.append("尚未建立DWG对象映射")
    return manifest, problems


def readiness(project):
    from .core import validate
    problems = validate(project)
    if not project.get("external_checks", "").strip():
        problems.append("请记录外部稳定与变形验算结论")
    if not project.get("template_manifest"):
        problems.append("尚未接入审核后的Excel与DWG模板映射")
    else:
        try:
            _, template_problems = check_manifest(project["template_manifest"])
            problems.extend(template_problems)
        except (OSError, ValueError, TypeError, AttributeError) as exc:
            problems.append(f"模板映射无法读取：{exc}")
    # Real geometry operations and engineering rules must be calibrated with supplied templates.
    problems.append("当前为模板接入工作台；自动选型与正式成果生成尚未完成基准验证")
    return problems
