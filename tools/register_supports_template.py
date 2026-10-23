"""Register the engineer-selected AA source after template review."""
import argparse,json
from pathlib import Path
from openpyxl import load_workbook
from els.templates import sha256
from els.supports import ROOT,BLINDING_FILE
from els.connections import copy_source
from tempfile import TemporaryDirectory

def main():
 parser=argparse.ArgumentParser();parser.add_argument('--source',type=Path,default=BLINDING_FILE);args=parser.parse_args()
 source=args.source.resolve()
 with TemporaryDirectory(prefix='els-aa-register-') as folder:
  snapshot=Path(folder)/'aa.xlsx';copy_source(source,snapshot);fingerprint=sha256(snapshot)
  book=load_workbook(snapshot,read_only=True,data_only=False,keep_links=False)
  try:
   sheet=book['Blinding'];formulas={a:sheet[a].value for a in ('D13','D18','C22','C26','G15','G16')}
   if any(not isinstance(f,str) or not f.startswith('=') for f in formulas.values()):raise ValueError('Missing Blinding formulas')
   for name in ('Connection check','Bracket','CS1 ','W1','D-Strut&Wailing Summary','Section Table UC'):
    if name not in book.sheetnames:raise ValueError('Missing AA worksheet: '+name)
  finally:book.close()
 config={'schema_version':2,'blinding_file':str(source),'blinding_sha256':fingerprint,'connection_file':str(source),'connection_sha256':fingerprint,'expected_formulas':formulas,'approved_source':'Engineer-selected AA Connection check, Bracket and Blinding worksheets; controlled load links and comparison repairs'}
 destination=ROOT/'templates/supports-structure.local.json';destination.write_text(json.dumps(config,ensure_ascii=False,indent=2),encoding='utf-8');print(destination)
if __name__=='__main__':main()
